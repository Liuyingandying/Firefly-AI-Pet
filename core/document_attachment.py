"""Companion Document Attachment v1 — parsing, chunking and retrieval.

Documents are parsed LOCALLY and only relevant excerpts ever leave the
machine. No raw document bytes, no absolute path and no full extracted text
are persisted anywhere: the conversation history keeps only a safe
``[Document attachment: <name>]`` placeholder plus the user's question.

Pipeline for one turn:
    user file -> local parser -> DocumentContext (sections) -> chunks
    -> direct page/slide lookup OR BM25 retrieval -> text excerpts
    -> text provider (core.ai_router) -> Companion answer

PDF parsing reuses the existing REAL PASS pipeline
(``core.pdf_processor.PdfProcessor``: PyMuPDF native text layer + RapidOCR
fallback for scanned pages). DOCX / PPTX / XLSX use python-docx /
python-pptx / openpyxl; TXT / MD / CSV use the standard library.
"""

from __future__ import annotations

import csv
import io
import logging
import math
import re
import time
from dataclasses import dataclass, field
from typing import Any, Sequence

log = logging.getLogger(__name__)

# ---------------------------------------------------------------- models


@dataclass
class DocumentSection:
    """One addressable slice of a document (page / slide / sheet / section)."""

    index: int
    label: str  # "Page 12" / "Slide 8" / "Sheet: Data" / "Section 3"
    text: str


@dataclass
class DocumentContext:
    """Fully parsed, in-memory document content. Never persisted."""

    filename: str
    kind: str
    sections: list[DocumentSection] = field(default_factory=list)
    total_characters: int = 0
    page_count: int | None = None
    slide_count: int | None = None
    sheet_count: int | None = None
    truncated: bool = False
    parse_warnings: list[str] = field(default_factory=list)


@dataclass
class DocumentChunk:
    """A retrieval unit tagged with its source label."""

    index: int
    label: str
    text: str
    section_index: int


# ---------------------------------------------------------------- limits

MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_TEXT_BYTES = 20 * 1024 * 1024
MAX_XLSX_BYTES = 30 * 1024 * 1024

DOCUMENT_KIND_LIMITS: dict[str, int] = {
    "pdf": MAX_DOCUMENT_BYTES,
    "docx": MAX_DOCUMENT_BYTES,
    "pptx": MAX_DOCUMENT_BYTES,
    "txt": MAX_TEXT_BYTES,
    "md": MAX_TEXT_BYTES,
    "csv": MAX_TEXT_BYTES,
    "xlsx": MAX_XLSX_BYTES,
}

MAX_PDF_PAGES = 500
CHUNK_SIZE = 1200          # ~chars per chunk (task: 1000-2000)
CHUNK_OVERLAP = 100

# Summary latency v2 — fewer remote calls via local structure compression.
#
# Path selection is based on the size of the LOCALLY COMPRESSED summary input
# (not the raw document): small -> 1 call, medium -> 2 map groups + 1 reduce
# (3 calls), large -> bounded map/reduce (<= SUMMARY_MAX_MAP_GROUPS + 1).
# The values are conservative character budgets; Chinese text is roughly one
# token per character, and we reserve headroom for system prompt, question,
# persona and the final answer.
SUMMARY_DIRECT_CHAR_BUDGET = 18000     # LEVEL 1: whole outline in one call
SUMMARY_MEDIUM_CHAR_BUDGET = 36000     # LEVEL 2: two map groups -> <=3 calls
SUMMARY_GROUP_CHAR_BUDGET = 18000      # per-map-group budget (LEVEL 3)
SUMMARY_MAX_MAP_GROUPS = 5             # hard cap on map groups (LEVEL 3)
SUMMARY_PRIORITY_SECTION_BUDGET = 1800  # per academic heading section
SUMMARY_NORMAL_SECTION_BUDGET = 900     # per ordinary section
SUMMARY_REFERENCE_BUDGET = 150          # references keep metadata only
SUMMARY_LEAD_FRACTION = 0.35
SUMMARY_TAIL_FRACTION = 0.20
SUMMARY_FREQUENT_TERMS = 12

RETRIEVAL_TOP_K = 6

MAX_XLSX_SHEET_ROWS = 2000
MAX_XLSX_CELLS = 100_000
MAX_CSV_ROWS = 5000

KIND_LABELS = {
    "pdf": "PDF", "docx": "DOCX", "pptx": "PPTX",
    "txt": "TXT", "md": "Markdown", "xlsx": "XLSX", "csv": "CSV",
}

# ---------------------------------------------------------------- errors


class DocumentParseError(Exception):
    """User-facing document parse failure (corrupt / encrypted / unsupported)."""


class EncryptedPdfError(DocumentParseError):
    """The PDF requires a password."""


# ---------------------------------------------------------------- dispatch


def parse_document_bytes(data: bytes, kind: str, filename: str) -> DocumentContext:
    """Parse one in-memory document into a :class:`DocumentContext`."""
    if kind == "pdf":
        return _parse_pdf(data, filename)
    if kind == "docx":
        return _parse_docx(data, filename)
    if kind == "pptx":
        return _parse_pptx(data, filename)
    if kind == "txt" or kind == "md":
        return _parse_text(data, filename, kind)
    if kind == "csv":
        return _parse_csv(data, filename)
    if kind == "xlsx":
        return _parse_xlsx(data, filename)
    raise DocumentParseError(f"unsupported document kind: {kind}")


# ---------------------------------------------------------------- PDF


def _parse_pdf(data: bytes, filename: str) -> DocumentContext:
    from core.pdf_processor import PdfEncryptedError, PdfProcessor

    try:
        result = PdfProcessor().process_stream(data, filename, max_pages=MAX_PDF_PAGES)
    except PdfEncryptedError as exc:
        raise EncryptedPdfError("PDF 需要密码") from exc
    except Exception as exc:
        raise DocumentParseError("无法读取 PDF") from exc

    sections = [
        DocumentSection(index=page.page_index + 1, label=f"Page {page.page_index + 1}", text=page.text)
        for page in result.pages
    ]
    ocr_errors = [err for err in result.errors if "OCR" in err or "ocr" in err.lower()]
    warnings = list(result.errors)
    if result.ocr_used:
        warnings.insert(0, "部分页面通过 OCR 识别")
    elif result.ocr_attempted:
        warnings.insert(0, "部分页面尝试 OCR 识别但未获得文本")
    context = DocumentContext(
        filename=filename,
        kind="pdf",
        sections=sections,
        total_characters=sum(len(s.text) for s in sections),
        page_count=len(sections),
        truncated=bool(result.errors and "truncated" in result.errors[-1]),
        parse_warnings=warnings,
    )
    return context


# ---------------------------------------------------------------- DOCX


def _parse_docx(data: bytes, filename: str) -> DocumentContext:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - python-docx is a dependency
        raise DocumentParseError("DOCX 解析库不可用") from exc

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("无法读取 DOCX") from exc

    sections: list[DocumentSection] = []
    para_index = 0
    table_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = (para.text or "").strip()
            if not text:
                continue
            para_index += 1
            label = f"Section {para_index}"
            sections.append(DocumentSection(index=para_index, label=label, text=text))
        elif child.tag == qn("w:tbl"):
            try:
                table = Table(child, doc)
            except Exception:
                continue
            table_index += 1
            rows: list[str] = []
            for row in table.rows:
                cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                rows.append(" | ".join(cells))
            label = f"Table {table_index}"
            sections.append(DocumentSection(
                index=len(sections) + 1, label=label, text=f"[{label}]\n" + "\n".join(rows)
            ))

    return DocumentContext(
        filename=filename,
        kind="docx",
        sections=sections,
        total_characters=sum(len(s.text) for s in sections),
    )


# ---------------------------------------------------------------- PPTX


def _parse_pptx(data: bytes, filename: str) -> DocumentContext:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - python-pptx is a dependency
        raise DocumentParseError("PPTX 解析库不可用") from exc

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("无法读取 PPTX") from exc

    sections: list[DocumentSection] = []
    for slide_index, slide in enumerate(prs.slides, 1):
        parts: list[str] = []
        title_shape = slide.shapes.title
        title_text = ""
        if title_shape is not None:
            title_text = (title_shape.text or "").strip()
            if title_text:
                parts.append(title_text)
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_text_frame", False):
                text = (shape.text_frame.text or "").strip()
                if text:
                    parts.append(text)
            elif getattr(shape, "has_table", False):
                rows: list[str] = []
                for row in shape.table.rows:
                    cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    parts.append("[Table]\n" + "\n".join(rows))
        try:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                parts.append(f"[Notes]\n{notes_text}")
        except Exception:
            pass
        text = "\n".join(p for p in parts if p)
        sections.append(DocumentSection(index=slide_index, label=f"Slide {slide_index}", text=text))

    return DocumentContext(
        filename=filename,
        kind="pptx",
        sections=sections,
        total_characters=sum(len(s.text) for s in sections),
        slide_count=len(sections),
    )


# ---------------------------------------------------------------- TXT / MD


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_text(data: bytes, filename: str, kind: str) -> DocumentContext:
    text = _decode_text(data)
    sections = [DocumentSection(index=1, label="Document", text=text)] if text.strip() else []
    return DocumentContext(
        filename=filename,
        kind=kind,
        sections=sections,
        total_characters=len(text),
    )


# ---------------------------------------------------------------- CSV


def _parse_csv(data: bytes, filename: str) -> DocumentContext:
    text = _decode_text(data)
    rows: list[str] = []
    truncated = False
    try:
        reader = csv.reader(io.StringIO(text))
        for index, row in enumerate(reader):
            if index >= MAX_CSV_ROWS:
                truncated = True
                break
            rows.append(" | ".join((cell or "").strip() for cell in row))
    except Exception as exc:
        raise DocumentParseError("无法读取 CSV") from exc
    sections = [DocumentSection(index=1, label="CSV", text="\n".join(rows))]
    return DocumentContext(
        filename=filename,
        kind="csv",
        sections=sections,
        total_characters=sum(len(r) for r in rows),
        truncated=truncated,
        parse_warnings=["内容超过上限已截断"] if truncated else [],
    )


# ---------------------------------------------------------------- XLSX


def _parse_xlsx(data: bytes, filename: str) -> DocumentContext:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl is a dependency
        raise DocumentParseError("XLSX 解析库不可用") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentParseError("无法读取 XLSX") from exc

    sections: list[DocumentSection] = []
    total_cells = 0
    truncated = False
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            row_count = 0
            for row in worksheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                rows.append(" | ".join(cells))
                row_count += 1
                total_cells += len(cells)
                if row_count >= MAX_XLSX_SHEET_ROWS or total_cells >= MAX_XLSX_CELLS:
                    truncated = True
                    break
            label = f"Sheet: {worksheet.title}"
            sections.append(DocumentSection(
                index=len(sections) + 1, label=label, text="\n".join(rows)
            ))
    finally:
        workbook.close()

    return DocumentContext(
        filename=filename,
        kind="xlsx",
        sections=sections,
        total_characters=sum(len(s.text) for s in sections),
        sheet_count=len(sections),
        truncated=truncated,
        parse_warnings=["表格内容超过上限已截断"] if truncated else [],
    )


# ---------------------------------------------------------------- chunking


def chunk_document(
    context: DocumentContext,
    *,
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[DocumentChunk]:
    """Split every section into source-labelled chunks of ~``chunk_size``."""
    chunks: list[DocumentChunk] = []
    for section in context.sections:
        text = (section.text or "").strip()
        if not text:
            continue
        start = 0
        guard = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            chunks.append(DocumentChunk(
                index=len(chunks),
                label=section.label,
                text=text[start:end],
                section_index=section.index,
            ))
            if end == len(text):
                break
            start = max(start + chunk_size - overlap, start + 1)
            guard += 1
            if guard > 2000:  # hard safety net for degenerate inputs
                break
    return chunks


# ---------------------------------------------------------------- retrieval


_WORD_RE = re.compile(r"[A-Za-z0-9_]+")
_STOP_WORDS = {
    "the", "a", "an", "of", "in", "on", "and", "or", "to", "for", "is",
    "are", "was", "were", "this", "that", "it", "with", "as", "at", "by",
    "from", "what", "how", "why", "do", "does", "did", "page", "slide",
}


def _tokenize(text: str) -> list[str]:
    tokens: list[str] = []
    for word in _WORD_RE.findall(text.lower()):
        if word not in _STOP_WORDS:
            tokens.append(word)
    # CJK unigrams carry the meaning for Chinese questions.
    for char in text:
        if "\u4e00" <= char <= "\u9fff":
            tokens.append(char)
    return tokens


def bm25_scores(
    query_terms: Sequence[str],
    chunk_terms: Sequence[Sequence[str]],
) -> list[float]:
    """Lightweight BM25 (k1=1.5, b=0.75) over in-memory chunks. No index."""
    n_chunks = len(chunk_terms)
    if n_chunks == 0 or not query_terms:
        return [0.0] * n_chunks
    doc_lengths = [max(len(terms), 1) for terms in chunk_terms]
    average_length = sum(doc_lengths) / n_chunks

    document_frequency: dict[str, int] = {}
    for terms in chunk_terms:
        for term in set(terms):
            document_frequency[term] = document_frequency.get(term, 0) + 1

    k1, b = 1.5, 0.75
    scores: list[float] = []
    idf_cache: dict[str, float] = {}
    for index, terms in enumerate(chunk_terms):
        term_counts: dict[str, int] = {}
        for term in terms:
            term_counts[term] = term_counts.get(term, 0) + 1
        score = 0.0
        for term in query_terms:
            if term not in term_counts:
                continue
            if term not in idf_cache:
                df = document_frequency.get(term, 0)
                idf_cache[term] = math.log(
                    1 + (n_chunks - df + 0.5) / (df + 0.5)
                ) if df else 0.0
            tf = term_counts[term]
            denom = tf + k1 * (1 - b + b * doc_lengths[index] / average_length)
            score += idf_cache[term] * (tf * (k1 + 1)) / max(denom, 1e-9)
        scores.append(score)
    return scores


def retrieve_chunks(question: str, chunks: list[DocumentChunk], top_k: int = RETRIEVAL_TOP_K) -> list[DocumentChunk]:
    """Return the most relevant chunks for a question (BM25 over chunks)."""
    if not chunks:
        return []
    query_terms = _tokenize(question)
    if not query_terms:
        return chunks[:top_k]
    scored = sorted(
        zip(chunks, bm25_scores(query_terms, [_tokenize(c.text) for c in chunks])),
        key=lambda pair: pair[1],
        reverse=True,
    )
    return [chunk for chunk, _score in scored[:top_k] if _score > 0] or chunks[:1]


# ---------------------------------------------------------------- direct lookup

_PAGE_CN_RE = re.compile(r"第\s*(\d+)\s*页")
_PAGE_EN_RE = re.compile(r"\bpage\s*(\d+)\b", re.IGNORECASE)
_SLIDE_RE = re.compile(r"\bslide\s*(\d+)\b", re.IGNORECASE)
_SHEET_RE = re.compile(r"sheet\s*[::：]?\s*[\"']?([^\"'\s，。]+)", re.IGNORECASE)


def direct_section_lookup(question: str, context: DocumentContext) -> list[DocumentSection]:
    """Deterministic page / slide / sheet lookup, no LLM intent classifier."""
    normalized = question.strip().lower()
    targets: list[str] = []

    match = _PAGE_CN_RE.search(question) or _PAGE_EN_RE.search(question)
    if match:
        number = int(match.group(1))
        if context.kind == "pptx":
            targets.append(f"Slide {number}")
        else:
            targets.append(f"Page {number}")

    match = _SLIDE_RE.search(normalized)
    if match and context.kind == "pptx":
        targets.append(f"Slide {int(match.group(1))}")

    match = _SHEET_RE.search(normalized)
    if match and context.kind == "xlsx":
        targets.append(f"Sheet: {match.group(1)}")

    if not targets:
        return []
    return [section for section in context.sections if section.label in targets]


def format_excerpts(sections_or_chunks: Sequence[Any], *, label_attr: str = "label", text_attr: str = "text") -> str:
    """Render retrieval units as the bounded SOURCE EXCERPTS block."""
    blocks = []
    for unit in sections_or_chunks:
        text = (getattr(unit, text_attr) or "").strip()
        if not text:
            continue
        blocks.append(f"[{getattr(unit, label_attr)}]\n{text}")
    return "\n\n".join(blocks)


# ---------------------------------------------------------------- summary intent


_SUMMARY_KEYWORDS = (
    "总结", "概述", "摘要", "概括", "要点", "主要内容", "整体逻辑",
    "讲了什么", "讲了些什么", "写论文笔记", "逐页讲", "详细总结",
    "summary", "summarize", "overview", "main points", "主要内容是什么",
)


_PAGE_OR_SLIDE_HINT_RE = re.compile(
    r"第\s*\d+\s*页|\bpage\s*\d+\b|\bslide\s*\d+\b", re.IGNORECASE
)


def is_summary_request(question: str) -> bool:
    lowered = (question or "").strip().lower()
    if not lowered:
        return False
    # "第 6 页讲了什么" is a page-specific question, never a full summary.
    if _PAGE_OR_SLIDE_HINT_RE.search(lowered):
        return False
    return any(keyword in lowered for keyword in _SUMMARY_KEYWORDS)


# ---------------------------------------------------------------- provider messages


DOCUMENT_PERSONA = """You are Firefly, the user's desktop companion.
Answer based only on the supplied document excerpts.
If the excerpts do not support the answer, say so.
Do not invent content.
Be warm and concise, normally 2-5 sentences.
Refer to pages/slides naturally when the content comes from a clear source."""

SUMMARY_MAP_PROMPT = (
    "你是 Firefly，用户的桌面伙伴。请阅读以下文档片段，用简洁要点概括其内容，"
    "保留重要细节与页/幻灯片来源。不要编造片段中没有的内容。\n\n"
)

SUMMARY_DIRECT_PROMPT = (
    "你是 Firefly，用户的桌面伙伴。请基于以下文档内容，用五点总结这份文档：\n"
    "1. 研究问题/目的\n"
    "2. 方法\n"
    "3. 数据/实验\n"
    "4. 主要结果\n"
    "5. 局限/意义\n"
    "只依据提供的文档内容；如果某项文档未明确提供，请明确说明“文档未明确提供”，"
    "不要臆造。内容来自哪些页/幻灯片请自然提及，不必逐句引用。"
    "保持简洁、温暖。\n\n"
)

SUMMARY_REDUCE_PROMPT = (
    "以下是同一文档各部分的分段摘要。请合并为一份完整、连贯的中文总结，"
    "并用五点组织：研究问题/目的、方法、数据/实验、主要结果、局限/意义。"
    "只依据所提供的分段摘要；某部分未提供时明确说明“文档未明确提供”，不要臆造。\n\n"
)


def build_qa_messages(
    filename: str,
    excerpts: str,
    question: str,
    *,
    persona: str = DOCUMENT_PERSONA,
) -> list[dict[str, str]]:
    """Bounded model payload: document name + relevant excerpts + question.

    Never contains raw file bytes, the absolute path, memory, or history.
    """
    user = (
        f"DOCUMENT:\n{filename}\n\n"
        f"SOURCE EXCERPTS:\n\n{excerpts}\n\n"
        f"USER QUESTION:\n{question}"
    )
    return [
        {"role": "system", "content": persona},
        {"role": "user", "content": user},
    ]


# ---------------------------------------------------------------- summary v2


@dataclass
class SummaryEntry:
    """One source-labelled entry of the locally compressed summary input."""

    label: str
    text: str


@dataclass
class SummaryInput:
    """Deterministic, locally compressed document outline. Never persisted."""

    filename: str
    kind: str
    entries: list[SummaryEntry] = field(default_factory=list)
    total_chars: int = 0
    references_downgraded: int = 0
    detected_headings: list[str] = field(default_factory=list)


_ACADEMIC_HEADINGS_EN: dict[str, tuple[str, ...]] = {
    "abstract": ("abstract",),
    "introduction": ("introduction",),
    "background": ("background",),
    "related_work": ("related work", "related works"),
    "method": ("method", "methods", "methodology", "approach"),
    "experiments": ("experiment", "experiments", "evaluation", "evaluations"),
    "results": ("results", "result"),
    "discussion": ("discussion",),
    "conclusion": ("conclusion", "conclusions", "summary"),
    "references": ("references", "reference", "bibliography"),
}

_ACADEMIC_HEADINGS_ZH: dict[str, tuple[str, ...]] = {
    "abstract": ("摘要",),
    "introduction": ("引言", "绪论", "前言"),
    "background": ("背景", "研究背景"),
    "related_work": ("相关工作", "国内外研究"),
    "method": ("方法", "研究方法", "模型设计", "方案设计"),
    "experiments": ("实验", "试验", "仿真", "评估"),
    "results": ("结果", "实验结果", "结果与分析"),
    "discussion": ("讨论",),
    "conclusion": ("结论", "总结", "结束语"),
    "references": ("参考文献", "参考资料"),
}

_REFERENCE_KINDS = frozenset({"references"})
_PRIORITY_KINDS = frozenset({
    "abstract", "introduction", "background", "related_work", "method",
    "experiments", "results", "discussion", "conclusion",
})

_RESULT_INDICATORS = (
    "accuracy", "result", "achieve", "performance", "improve", "compared",
    "outperform", "propose", "proposed", "show", "demonstrate",
    "准确率", "精度", "结果", "性能", "优于", "提升", "提出", "表明",
)


def _first_nonempty_line(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return ""


def _normalize_heading_line(line: str) -> str:
    return line.strip().lower().lstrip("0123456789.、·-—()（） \t:：\"'「」[]")


def _looks_like_references(text: str) -> bool:
    sample = text[:4000]
    bracket_lines = sum(
        1 for line in sample.splitlines() if re.match(r"^\s*\[\d+\]", line)
    )
    numbered_lines = sum(
        1 for line in sample.splitlines() if re.match(r"^\s*\d+\.\s", line)
    )
    return bracket_lines >= 3 or (numbered_lines >= 5 and len(sample) > 2000)


def detect_section_kind(text: str) -> str | None:
    """Classify a section by its first line / citation density (no ML).

    The longest matching heading keyword wins, so ``实验结果`` classifies as
    ``results`` rather than the shorter ``实验`` -> ``experiments``.
    """
    first = _first_nonempty_line(text)[:60]
    normalized = _normalize_heading_line(first)
    best_kind: str | None = None
    best_length = 0
    for kind, keywords in _ACADEMIC_HEADINGS_EN.items():
        for keyword in keywords:
            if (
                normalized == keyword
                or normalized.startswith(keyword + " ")
                or normalized.startswith(keyword + ":")
                or normalized.startswith(keyword + ".")
            ):
                if len(keyword) > best_length:
                    best_kind, best_length = kind, len(keyword)
    for kind, keywords in _ACADEMIC_HEADINGS_ZH.items():
        for keyword in keywords:
            if normalized.startswith(keyword):
                if len(keyword) > best_length:
                    best_kind, best_length = kind, len(keyword)
    if best_kind is not None:
        return best_kind
    if _looks_like_references(text):
        return "references"
    return None


def _dedupe_lines(text: str) -> str:
    seen: set[str] = set()
    kept: list[str] = []
    for line in text.splitlines():
        key = line.strip().lower()
        if not key:
            continue
        if re.fullmatch(r"[\s\d.\-–—|·]+", key):  # page numbers / separators
            continue
        if key in seen:
            continue
        seen.add(key)
        kept.append(line.strip())
    return "\n".join(kept)


def _split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[。！？.!?；;])\s*|\n+", text)
    return [part.strip() for part in parts if part.strip()]


def _frequent_terms(context: DocumentContext, top_n: int = SUMMARY_FREQUENT_TERMS) -> list[str]:
    counts: dict[str, int] = {}
    for section in context.sections:
        lowered = section.text.lower()
        for word in _WORD_RE.findall(lowered):
            if word in _STOP_WORDS or len(word) < 3:
                continue
            counts[word] = counts.get(word, 0) + 1
        for index in range(len(section.text) - 1):
            pair = section.text[index:index + 2]
            if "\u4e00" <= pair[0] <= "\u9fff" and "\u4e00" <= pair[1] <= "\u9fff":
                counts[pair] = counts.get(pair, 0) + 1
    ranked = sorted(counts.items(), key=lambda item: item[1], reverse=True)
    return [term for term, _count in ranked[:top_n]]


def _selective_sentences(text: str, budget: int, frequent_terms: list[str]) -> str:
    sentences = _split_sentences(text)
    scored: list[tuple[int, int, str]] = []  # (score desc, original index, sentence)
    for index, sentence in enumerate(sentences):
        score = 0
        if any(char.isdigit() for char in sentence):
            score += 1
        lowered = sentence.lower()
        for term in frequent_terms:
            if term and term.lower() in lowered:
                score += 1
        if any(indicator in lowered for indicator in _RESULT_INDICATORS):
            score += 1
        if score > 0:
            scored.append((-score, index, sentence))
    scored.sort()
    selected: list[tuple[int, str]] = []
    total = 0
    for _neg_score, index, sentence in scored:
        if total + len(sentence) > budget:
            continue
        selected.append((index, sentence))
        total += len(sentence)
    selected.sort(key=lambda item: item[0])
    return "\n".join(sentence for _index, sentence in selected)


def _extract_section_text(text: str, budget: int, frequent_terms: list[str]) -> str:
    text = _dedupe_lines(text)
    if len(text) <= budget:
        return text
    lead = text[: int(budget * SUMMARY_LEAD_FRACTION)]
    tail = text[-int(budget * SUMMARY_TAIL_FRACTION):]
    middle_budget = max(0, budget - len(lead) - len(tail))
    middle = _selective_sentences(text, middle_budget, frequent_terms)
    return "\n".join(part for part in (lead, middle, tail) if part.strip())


def build_summary_input(context: DocumentContext) -> SummaryInput:
    """Deterministic local compression into a source-labelled outline.

    Academic headings (EN + ZH) keep more text; references keep metadata only;
    repeated header/footer lines are dropped. The original DocumentContext is
    never modified — this produces a separate SummaryInput.
    """
    frequent = _frequent_terms(context)
    entries: list[SummaryEntry] = []
    detected: list[str] = []
    references_downgraded = 0
    for section in context.sections:
        text = (section.text or "").strip()
        if not text:
            continue
        kind = detect_section_kind(text)
        if kind == "references":
            entries.append(SummaryEntry(
                label=section.label, text=text[:SUMMARY_REFERENCE_BUDGET]
            ))
            references_downgraded += 1
            continue
        if kind is not None:
            detected.append(kind)
        budget = (
            SUMMARY_PRIORITY_SECTION_BUDGET
            if kind in _PRIORITY_KINDS
            else SUMMARY_NORMAL_SECTION_BUDGET
        )
        entries.append(SummaryEntry(
            label=section.label, text=_extract_section_text(text, budget, frequent)
        ))
    total = sum(len(entry.text) for entry in entries)
    return SummaryInput(
        filename=context.filename,
        kind=context.kind,
        entries=entries,
        total_chars=total,
        references_downgraded=references_downgraded,
        detected_headings=sorted(set(detected)),
    )


def format_summary_input(summary_input: SummaryInput) -> str:
    """Render the compressed outline as the SOURCE EXCERPTS block."""
    blocks = []
    for entry in summary_input.entries:
        text = entry.text.strip()
        if text:
            blocks.append(f"[{entry.label}]\n{text}")
    return "\n\n".join(blocks)


def _split_entries(entries: list[SummaryEntry], num_groups: int) -> list[list[SummaryEntry]]:
    total = sum(len(entry.text) for entry in entries)
    if num_groups <= 1 or total <= 0:
        return [entries] if entries else []
    target = max(1, math.ceil(total / num_groups))
    groups: list[list[SummaryEntry]] = []
    current: list[SummaryEntry] = []
    current_chars = 0
    for entry in entries:
        if current and current_chars + len(entry.text) > target:
            groups.append(current)
            current = []
            current_chars = 0
        current.append(entry)
        current_chars += len(entry.text)
    if current:
        groups.append(current)
    # Merge trailing leftover groups so we never exceed ``num_groups``.
    while len(groups) > num_groups and len(groups) >= 2:
        last = groups.pop()
        groups[-1].extend(last)
    return groups


def summarize_document(
    context: DocumentContext,
    chat: Any,
    *,
    cache: dict | None = None,
) -> tuple[str, int, dict]:
    """Full-document summary with few remote calls.

    Strategy (based on the locally compressed outline size):
      LEVEL 1 small   (<= SUMMARY_DIRECT_CHAR_BUDGET)         -> 1 call
      LEVEL 2 medium  (<= SUMMARY_MEDIUM_CHAR_BUDGET)         -> 2 groups -> 3 calls
      LEVEL 3 large   (> medium)  bounded map/reduce          -> <= MAX_GROUPS+1 calls

    ``cache`` (in-memory, session-scoped) reuses the compressed outline and
    the map-group summaries so a follow-up summary question only re-runs the
    cheap final synthesis. Returns ``(answer, remote_calls, stats)`` where
    stats = {summary_input_chars, groups, level, references_downgraded,
             detected_headings, local_prepare_ms}.
    """
    store = cache if isinstance(cache, dict) else {}
    prepared = time.perf_counter()
    summary_input = store.get("input")
    if summary_input is None:
        summary_input = build_summary_input(context)
        store["input"] = summary_input
    local_prepare_ms = (time.perf_counter() - prepared) * 1000

    total = summary_input.total_chars
    if total <= SUMMARY_DIRECT_CHAR_BUDGET:
        excerpts = format_summary_input(summary_input)
        messages = build_qa_messages(
            context.filename, excerpts, SUMMARY_DIRECT_PROMPT
        )
        answer = _chat_content(chat(messages, temperature=0.3))
        stats = {
            "summary_input_chars": total,
            "groups": 1,
            "level": "small",
            "references_downgraded": summary_input.references_downgraded,
            "detected_headings": summary_input.detected_headings,
            "local_prepare_ms": round(local_prepare_ms, 1),
        }
        return answer, 1, stats

    if total <= SUMMARY_MEDIUM_CHAR_BUDGET:
        groups = _split_entries(summary_input.entries, 2)
        level = "medium"
    else:
        num_groups = max(
            2, min(SUMMARY_MAX_MAP_GROUPS, math.ceil(total / SUMMARY_GROUP_CHAR_BUDGET))
        )
        groups = _split_entries(summary_input.entries, num_groups)
        level = "large"

    group_summaries = store.get("groups")
    map_calls_this_invocation = 0
    if group_summaries is None:
        group_summaries = []
        for group in groups:
            excerpts = format_summary_input(
                SummaryInput(
                    filename=summary_input.filename,
                    kind=summary_input.kind,
                    entries=group,
                    total_chars=sum(len(entry.text) for entry in group),
                )
            )
            messages = [
                {"role": "system", "content": DOCUMENT_PERSONA},
                {"role": "user", "content": SUMMARY_MAP_PROMPT + excerpts},
            ]
            group_summaries.append(_chat_content(chat(messages, temperature=0.3)))
            map_calls_this_invocation += 1
        store["groups"] = group_summaries

    combined = "\n\n".join(group_summaries)
    final_messages = [
        {"role": "system", "content": DOCUMENT_PERSONA},
        {"role": "user", "content": SUMMARY_REDUCE_PROMPT + combined},
    ]
    final = _chat_content(chat(final_messages, temperature=0.3))
    stats = {
        "summary_input_chars": total,
        "groups": len(group_summaries),
        "level": level,
        "references_downgraded": summary_input.references_downgraded,
        "detected_headings": summary_input.detected_headings,
        "local_prepare_ms": round(local_prepare_ms, 1),
    }
    # Count only the calls this invocation actually makes: cached group
    # summaries do not re-run the map stage.
    return final, map_calls_this_invocation + 1, stats


def _chat_content(response: Any) -> str:
    """Extract the assistant text from an OpenAI-compatible completion dict."""
    if not isinstance(response, dict):
        return ""
    choices = response.get("choices")
    if not isinstance(choices, list) or not choices:
        return ""
    message = choices[0].get("message") if isinstance(choices[0], dict) else None
    content = message.get("content") if isinstance(message, dict) else None
    return content.strip() if isinstance(content, str) else ""


__all__ = [
    "DocumentSection",
    "DocumentContext",
    "DocumentChunk",
    "DocumentParseError",
    "EncryptedPdfError",
    "parse_document_bytes",
    "chunk_document",
    "retrieve_chunks",
    "bm25_scores",
    "direct_section_lookup",
    "format_excerpts",
    "is_summary_request",
    "build_qa_messages",
    "SummaryInput",
    "SummaryEntry",
    "build_summary_input",
    "format_summary_input",
    "detect_section_kind",
    "summarize_document",
    "DOCUMENT_PERSONA",
    "DOCUMENT_KIND_LIMITS",
    "MAX_PDF_PAGES",
    "SUMMARY_DIRECT_CHAR_BUDGET",
    "SUMMARY_MEDIUM_CHAR_BUDGET",
    "SUMMARY_GROUP_CHAR_BUDGET",
    "SUMMARY_MAX_MAP_GROUPS",
]
