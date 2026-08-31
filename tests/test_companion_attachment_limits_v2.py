"""Companion Attachment Limits v2 tests (A-AG).

Two-layer safety: Layer 1 = raw file entry limits (PDF/DOCX/PPTX 150MB,
XLSX/CSV/TXT/MD 100MB, images 50MB); Layer 2 = parse resource budgets
(pages / slides / chars / rows / cells / pixels / ZIP metadata). No real
150MB fixtures are created — sizes come from mocked stat / stream / ZIP
metadata. DocumentContext schema, Summary v2, BM25 and direct page/slide
lookup are untouched.
"""

from __future__ import annotations

import io
import json
import math
import time
from pathlib import Path

import pytest
from PySide6.QtWidgets import QApplication

from core.document_attachment import (
    MAX_DOCUMENT_EXTRACTED_CHARS,
    MAX_PDF_OCR_PAGES,
    MAX_PDF_PAGES,
    MAX_PPTX_SLIDES,
    MAX_TEXT_EXTRACTED_CHARS,
    MAX_XLSX_ROWS_PER_SHEET,
    MAX_XLSX_TOTAL_CELLS,
    ZIP_BOMB_MESSAGE,
    DocumentContext,
    DocumentSection,
    audit_zip_container,
    build_summary_input,
    direct_section_lookup,
    parse_document_bytes,
    retrieve_chunks,
)
from core.pdf_processor import PdfPage, PdfResult
from ui import companion_attachment as attachment_mod
from ui.companion_attachment import (
    MAX_IMAGE_PIXELS,
    AttachmentError,
    DocumentAttachment,
    decode_attachment_bytes,
    load_document_file,
)
from ui.companion_chat_window import CompanionChatWindow
from ui.character_conversation_runner import CharacterConversationRunner


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


class _FakeStat:
    def __init__(self, size):
        self.st_size = size


def _fake_size(monkeypatch, path: Path, size: int) -> None:
    real_stat = Path.stat

    def fake_stat(self):
        if Path(self) == Path(path):
            return _FakeStat(size)
        return real_stat(self)

    monkeypatch.setattr(Path, "stat", fake_stat)


def _mb(value: float) -> int:
    return int(value * 1024 * 1024)


# ---------------------------------------------------------------- fixtures


def _tiny_pdf_bytes() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "limits test content")
    data = doc.tobytes()
    doc.close()
    return data


def _n_page_pdf(n: int) -> bytes:
    import fitz

    doc = fitz.open()
    for index in range(n):
        page = doc.new_page()
        page.insert_text((72, 72), f"page {index + 1} content")
    data = doc.tobytes()
    doc.close()
    return data


def _tiny_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("a" * 2000)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _tiny_pptx_bytes(slides: int = 5) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for index in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {index + 1}"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: int = 150) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for index in range(rows):
        ws.append([f"r{index}", index])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _ready_doc(kind: str, data: bytes, name: str) -> DocumentAttachment:
    attachment = DocumentAttachment(
        display_name=name, kind=kind, original_size=len(data), source_bytes=data
    )
    attachment.mark_ready(parse_document_bytes(data, kind, name))
    return attachment


# ================================================== A-J entry layer (fake stat)


@pytest.mark.parametrize(
    "name,content,kind,size_mb,accepted",
    [
        ("big.pdf", None, "pdf", 64.4, True),   # A
        ("big.pdf", None, "pdf", 149.0, True),  # B
        ("big.pdf", None, "pdf", 151.0, False),  # C
        ("deck.pptx", None, "pptx", 149.0, True),  # D
        ("deck.pptx", None, "pptx", 151.0, False),  # E
        ("report.docx", None, "docx", 149.0, True),  # F
        ("data.xlsx", None, "xlsx", 101.0, False),  # G
        ("data.xlsx", None, "xlsx", 99.0, True),   # H
        ("table.csv", None, "csv", 99.0, True),   # I
        ("notes.txt", None, "txt", 99.0, True),   # J
    ],
    ids=["a_pdf_64mb", "b_pdf_149mb", "c_pdf_151mb", "d_pptx_149mb", "e_pptx_151mb",
         "f_docx_149mb", "g_xlsx_101mb", "h_xlsx_99mb", "i_csv_99mb", "j_txt_99mb"],
)
def test_entry_layer_size_acceptance(tmp_path, monkeypatch, name, content, kind, size_mb, accepted):
    path = tmp_path / name
    if kind == "pdf":
        path.write_bytes(_tiny_pdf_bytes())
    elif kind == "pptx":
        path.write_bytes(_tiny_pptx_bytes(2))
    elif kind == "docx":
        path.write_bytes(_tiny_docx_bytes())
    elif kind == "xlsx":
        path.write_bytes(_xlsx_bytes(3))
    elif kind == "csv":
        path.write_bytes(b"a,b\n1,2\n")
    else:
        path.write_bytes("hello".encode("utf-8"))
    _fake_size(monkeypatch, path, _mb(size_mb))
    if accepted:
        attachment = load_document_file(path)
        assert attachment.kind == kind
        assert attachment.original_size == _mb(size_mb)
    else:
        with pytest.raises(AttachmentError):
            load_document_file(path)


def test_image_entry_limit_50mb(tmp_path, monkeypatch):
    from PIL import Image as PILImage

    path = tmp_path / "photo.png"
    buffer = io.BytesIO()
    PILImage.new("RGB", (32, 32), (10, 20, 30)).save(buffer, "PNG")
    path.write_bytes(buffer.getvalue())
    _fake_size(monkeypatch, path, _mb(49.9))
    attachment = attachment_mod.load_attachment_file(path)
    assert attachment is not None
    _fake_size(monkeypatch, path, _mb(50.1))
    with pytest.raises(AttachmentError) as exc:
        attachment_mod.load_attachment_file(path)
    assert "50 MB" in str(exc.value)


# ================================================== K-M PDF budgets


def test_k_pdf_800_pages_native_accepted():
    context = parse_document_bytes(_n_page_pdf(800), "pdf", "big.pdf")
    assert context.page_count == 800
    assert context.truncated is False


def test_l_pdf_over_page_budget_truncated(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_PDF_PAGES", 10)
    context = parse_document_bytes(_n_page_pdf(15), "pdf", "big.pdf")
    assert context.truncated is True
    assert context.page_count == 10


def test_m_pdf_ocr_page_budget_passed_and_surfaced(monkeypatch):
    calls: dict = {}

    class FakeProcessor:
        def process_stream(self, data, display_name, *, max_pages=None, max_ocr_pages=None,
                           keep_page_images=True):
            calls["max_pages"] = max_pages
            calls["max_ocr_pages"] = max_ocr_pages
            calls["keep_page_images"] = keep_page_images
            pages = [
                PdfPage(page_index=i, text="", has_text_layer=False,
                        extraction_source="empty", ocr_attempted=i < 2)
                for i in range(3)
            ]
            return PdfResult(
                file_path="scan.pdf", total_pages=3, pages=pages,
                ocr_attempted=True, errors=["page 3: OCR skipped (page budget)"],
            )

    monkeypatch.setattr("core.pdf_processor.PdfProcessor", FakeProcessor)
    context = parse_document_bytes(b"fake", "pdf", "scan.pdf")
    assert calls["max_pages"] == 1000
    assert calls["max_ocr_pages"] == 200
    assert calls["keep_page_images"] is False
    assert any("OCR" in warning for warning in context.parse_warnings)
    assert any("未执行 OCR" in warning for warning in context.parse_warnings)


# ================================================== N-S extracted budgets


def test_n_docx_over_char_budget_truncated(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_DOCUMENT_EXTRACTED_CHARS", 500)
    context = parse_document_bytes(_tiny_docx_bytes(), "docx", "big.docx")
    assert context.truncated is True
    assert context.total_characters <= 500


def test_o_pptx_over_slide_budget_truncated(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_PPTX_SLIDES", 3)
    context = parse_document_bytes(_tiny_pptx_bytes(5), "pptx", "big.pptx")
    assert context.truncated is True
    assert context.slide_count == 3


def test_p_xlsx_over_rows_per_sheet_truncated(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_XLSX_ROWS_PER_SHEET", 100)
    context = parse_document_bytes(_xlsx_bytes(150), "xlsx", "big.xlsx")
    assert context.truncated is True


def test_q_xlsx_total_cells_bounded(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_XLSX_TOTAL_CELLS", 60)
    context = parse_document_bytes(_xlsx_bytes(50), "xlsx", "big.xlsx")
    assert context.truncated is True


def test_r_csv_rows_bounded(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_CSV_ROWS", 100)
    big_csv = ("a,b\n" + "x,y\n" * 150).encode("utf-8")
    context = parse_document_bytes(big_csv, "csv", "big.csv")
    assert context.truncated is True


def test_s_txt_chars_bounded(monkeypatch):
    monkeypatch.setattr("core.document_attachment.MAX_TEXT_EXTRACTED_CHARS", 1000)
    context = parse_document_bytes(("内容" * 2000).encode("utf-8"), "txt", "big.txt")
    assert context.truncated is True
    assert context.total_characters <= 1000


# ================================================== T/U image pixels


def test_t_extreme_pixels_downscaled_safely(monkeypatch):
    monkeypatch.setattr(attachment_mod, "MAX_IMAGE_PIXELS", 100_000)
    # Decode a 400x400 PNG -> over the 100k budget -> downscaled.
    from PIL import Image as PILImage

    buffer = io.BytesIO()
    PILImage.new("RGB", (400, 400), (1, 2, 3)).save(buffer, "PNG")
    attachment = decode_attachment_bytes(buffer.getvalue(), "big.png")
    pixels = attachment.image.width() * attachment.image.height()
    assert pixels <= 100_000
    assert abs(attachment.image.width() - attachment.image.height()) <= 1


def test_u_under_budget_image_kept(monkeypatch):
    monkeypatch.setattr(attachment_mod, "MAX_IMAGE_PIXELS", 10_000_000)
    from PIL import Image as PILImage

    buffer = io.BytesIO()
    PILImage.new("RGB", (64, 32), (1, 2, 3)).save(buffer, "PNG")
    attachment = decode_attachment_bytes(buffer.getvalue(), "small.png")
    assert (attachment.image.width(), attachment.image.height()) == (64, 32)


# ================================================== V-X ZIP bomb audit


def _fake_zipinfo(file_size: int, compress_size: int):
    return type("ZipInfo", (), {"file_size": file_size, "compress_size": compress_size})


def test_v_zip_uncompressed_over_1gb_rejected():
    infolist = [_fake_zipinfo(1_500_000_000, 100_000)]
    with pytest.raises(Exception) as exc:
        audit_zip_container(b"", infolist=infolist)
    assert ZIP_BOMB_MESSAGE in str(exc.value)


def test_w_zip_entries_over_20000_rejected():
    infolist = [_fake_zipinfo(100, 100) for _ in range(20_001)]
    with pytest.raises(Exception) as exc:
        audit_zip_container(b"", infolist=infolist)
    assert ZIP_BOMB_MESSAGE in str(exc.value)


def test_x_zip_compression_ratio_over_100_rejected():
    infolist = [_fake_zipinfo(10_000_000, 10_000)]
    with pytest.raises(Exception) as exc:
        audit_zip_container(b"", infolist=infolist)
    assert ZIP_BOMB_MESSAGE in str(exc.value)


def test_zip_audit_runs_on_real_docx():
    # A legitimate docx passes the audit and parses normally.
    context = parse_document_bytes(_tiny_docx_bytes(), "docx", "ok.docx")
    assert context.kind == "docx"


# ================================================== Y stat-before-read


def test_y_oversized_rejected_before_read(tmp_path, monkeypatch):
    path = tmp_path / "big.pdf"
    path.write_bytes(_tiny_pdf_bytes())
    _fake_size(monkeypatch, path, _mb(151))
    reads: list = []
    real_read = Path.read_bytes

    def fake_read(self, *args, **kwargs):
        reads.append(str(self))
        return real_read(self, *args, **kwargs)

    monkeypatch.setattr(Path, "read_bytes", fake_read)
    with pytest.raises(AttachmentError):
        load_document_file(path)
    assert reads == []                       # never read the oversized file


def test_y2_text_entries_read_bounded(tmp_path, monkeypatch):
    # A 99MB txt accepted at entry layer still reads only a bounded prefix.
    path = tmp_path / "huge.txt"
    path.write_bytes("x".encode("utf-8"))
    _fake_size(monkeypatch, path, _mb(99))
    reads: list = []
    real_read_bytes = Path.read_bytes
    real_open = open

    monkeypatch.setattr(Path, "read_bytes", lambda self: reads.append("read_bytes") or b"")
    attachment = load_document_file(path)
    assert attachment is not None
    # _read_bounded uses open(); Path.read_bytes is never used for txt.
    assert reads == []


# ================================================== Z parse off UI thread


def test_z_large_parse_happens_off_ui_thread(tmp_path, qapp):
    path = tmp_path / "paper.pdf"
    path.write_bytes(_tiny_pdf_bytes())
    window = CompanionChatWindow(
        runner=CharacterConversationRunner(
            runtime=type("R", (), {"chat": lambda *a, **k: {"choices": [{"message": {"role": "assistant", "content": "ok"}}]}})()
        )
    )
    window._set_attachment(load_document_file(path))
    assert window._pending_attachment.parse_state == "parsing"  # async, not blocked
    assert "Parsing" in window._chip.status_label.text()
    deadline = time.monotonic() + 15.0
    while time.monotonic() < deadline and window._pending_attachment.parse_state not in ("ready", "error"):
        QApplication.processEvents()
        time.sleep(0.01)
    QApplication.processEvents()
    assert window._pending_attachment.parse_state == "ready"
    # The UI thread was free throughout: we could call window methods mid-parse.
    assert window.pick_button.isEnabled()


def test_z2_large_file_chip_shows_large_parsing(tmp_path, qapp):
    attachment = DocumentAttachment(
        display_name="big.pdf", kind="pdf", original_size=_mb(64.4), source_bytes=b""
    )
    window = CompanionChatWindow()
    window._set_attachment(attachment)
    assert "正在解析大型文件" in window._chip.status_label.text()


# ================================================== AA-AD unchanged core


def test_aa_document_context_schema_unchanged():
    section = DocumentSection(index=1, label="Page 1", text="text")
    assert section.index == 1 and section.label == "Page 1" and section.text == "text"
    context = DocumentContext(
        filename="a.pdf", kind="pdf", sections=[section],
        total_characters=4, page_count=1, truncated=False, parse_warnings=[],
    )
    assert context.filename == "a.pdf"


def test_ab_summary_v2_unchanged():
    from core.document_attachment import build_summary_input

    context = parse_document_bytes(("内容" * 100).encode("utf-8"), "txt", "s.txt")
    summary_input = build_summary_input(context)
    assert summary_input.total_chars > 0


def test_ac_bm25_qa_unchanged():
    chunks = []
    for index in range(1, 6):
        chunks.append(type("C", (), {"label": f"Page {index}", "text": f"topic {index} details", "section_index": index, "index": index})())
    top = retrieve_chunks("topic 3", chunks, top_k=2)
    assert any(chunk.label == "Page 3" for chunk in top)


def test_ad_page_direct_lookup_unchanged():
    sections = [DocumentSection(index=i, label=f"Page {i}", text=f"c{i}") for i in range(1, 15)]
    context = DocumentContext(filename="p.pdf", kind="pdf", sections=sections, total_characters=10)
    hit = direct_section_lookup("第 12 页讲了什么", context)
    assert [s.label for s in hit] == ["Page 12"]


# ================================================== AE-AG privacy & release


class _EchoChat:
    def __init__(self):
        self.calls = []

    def __call__(self, messages, temperature=0.2):
        self.calls.append(messages)
        return {"choices": [{"message": {"role": "assistant", "content": "ok"}}]}


def test_ae_raw_document_never_uploaded():
    chat = _EchoChat()
    runner = CharacterConversationRunner(
        runtime=type("R", (), {"chat": lambda *a, **k: {"choices": [{"message": {"role": "assistant", "content": "ok"}}]}})(),
        document_chat_handler=chat,
    )
    runner.perform_with_document(
        "方法", _ready_doc("txt", ("内容" * 50).encode("utf-8"), "doc.txt")
    )
    blob = json.dumps(chat.calls, ensure_ascii=False)
    assert "data:" not in blob and "base64" not in blob


def test_af_memory_history_privacy_unchanged():
    chat = _EchoChat()
    runner = CharacterConversationRunner(
        runtime=type("R", (), {"chat": lambda *a, **k: {"choices": [{"message": {"role": "assistant", "content": "ok"}}]}})(),
        document_chat_handler=chat,
    )
    runner.perform_with_document("方法", _ready_doc("txt", b"content", "d.txt"))
    history = runner.history
    assert "[Document attachment: d.txt]" in history[0]["content"]
    assert "base64" not in json.dumps(history)


def test_ag_raw_bytes_released_after_parse():
    data = ("内容" * 20).encode("utf-8")
    attachment = DocumentAttachment(
        display_name="d.txt", kind="txt", original_size=len(data), source_bytes=data
    )
    assert attachment.take_source_bytes() == data
    attachment.mark_ready(parse_document_bytes(data, "txt", "d.txt"))
    assert attachment.take_source_bytes() == b""      # raw bytes released
    assert attachment.ready
